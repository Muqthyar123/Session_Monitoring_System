import io
import pytest
from datetime import datetime, timezone
import zoneinfo
import openpyxl
from bson import ObjectId
from app.core.config import settings
from app.db.mongodb import get_database
from app.schemas.planned_absence import (
    PlannedAbsenceCreateRequest,
    PlannedAbsenceUpdateRequest,
    PlannedAbsenceCancelRequest,
)
from app.schemas.mentor_mapping import MentorMappingConfirmRequest, MentorMappingRow
from app.services.planned_absence_service import (
    create_planned_absence,
    get_planned_absences,
    update_planned_absence,
    cancel_planned_absence,
    get_active_planned_absence_for_student,
    batch_get_active_planned_absences_map,
)
from app.services.mentor_mapping_service import (
    generate_mentor_mapping_excel_template,
    preview_mentor_mappings,
    commit_mentor_mappings,
    get_mentor_mappings,
    delete_mentor_mapping,
    get_mentor_assigned_student_ids,
    is_student_assigned_to_mentor,
)
from app.services.mentor_service import (
    get_mentor_dashboard_data,
    get_mentor_students_list,
    get_mentor_scoped_absentees,
)


@pytest.mark.asyncio
async def test_mentor_mapping_excel_template_and_preview():
    db = get_database()
    now_utc = datetime.now(timezone.utc)

    # 1. Template generation
    template_bytes = generate_mentor_mapping_excel_template()
    assert len(template_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Mentor Name"
    assert ws.cell(row=1, column=2).value == "Year"
    assert ws.cell(row=1, column=3).value == "Section"
    assert ws.cell(row=1, column=4).value == "Starting Serial Number"
    assert ws.cell(row=1, column=5).value == "Ending Serial Number"

    # Seed mentor and students
    mentor_res = await db.users.insert_one({
        "name": "Dr. Test Mentor Mapping",
        "email": "test.mentor.map@college.edu",
        "role": "MENTOR",
        "created_at": now_utc,
    })
    mentor_id = str(mentor_res.inserted_id)

    # Seed 10 students in 2nd Year Section A
    students_to_insert = []
    for i in range(1, 11):
        students_to_insert.append({
            "name": f"Student {i:02d}",
            "roll_number": f"24TEST05{i:02d}",
            "year": "2nd Year",
            "section": "A",
            "created_at": now_utc,
            "updated_at": now_utc,
        })
    await db.students.insert_many(students_to_insert)

    # 2. Create custom workbook with:
    # Row 1: Dr. Test Mentor Mapping | 2nd Year | A | 1 | 5  (Valid)
    # Row 2: Dr. Test Mentor Mapping | 2nd Year | A | 6 | 10 (Valid)
    wb_test = openpyxl.Workbook()
    ws_test = wb_test.active
    ws_test.append(["Mentor Name", "Year", "Section", "Starting Serial Number", "Ending Serial Number"])
    ws_test.append(["Dr. Test Mentor Mapping", "2nd Year", "A", "1", "5"])
    ws_test.append(["Dr. Test Mentor Mapping", "2nd Year", "A", "6", "10"])

    buf = io.BytesIO()
    wb_test.save(buf)
    file_bytes = buf.getvalue()

    preview = await preview_mentor_mappings(file_bytes, "test_mapping.xlsx")
    assert preview.total_rows == 2
    assert preview.valid_rows == 2
    assert preview.invalid_rows == 0
    assert preview.students_to_assign == 10

    # 3. Commit mappings
    confirm_req = MentorMappingConfirmRequest(
        mode="ADD_UPDATE",
        rows=[r.model_dump() for r in preview.rows],
    )
    commit_res = await commit_mentor_mappings(confirm_req, {"_id": ObjectId(mentor_id), "email": "admin@college.edu"})
    assert commit_res["success"] is True

    # Check student assignments
    assigned_ids = await get_mentor_assigned_student_ids({"_id": ObjectId(mentor_id), "role": "MENTOR", "name": "Dr. Test Mentor Mapping"})
    assert assigned_ids is not None
    assert len(assigned_ids) == 10

    # Check mapping list
    mappings = await get_mentor_mappings(year="2nd Year", section="A")
    assert len(mappings) >= 2

    # Clean up
    await db.users.delete_one({"_id": ObjectId(mentor_id)})
    await db.students.delete_many({"roll_number": {"$regex": "^24TEST05"}})
    await db.mentor_student_mappings.delete_many({"mentor_id": mentor_id})


@pytest.mark.asyncio
async def test_planned_absence_workflow_and_rbac():
    db = get_database()
    now_utc = datetime.now(timezone.utc)
    today_ist = datetime.now(zoneinfo.ZoneInfo("Asia/Kolkata")).strftime("%Y-%m-%d")

    # 1. Create Mentor and 2 Students (1 assigned, 1 unassigned)
    m_res = await db.users.insert_one({
        "name": "Prof. Planned Absence Mentor",
        "email": "mentor.pa@college.edu",
        "role": "MENTOR",
        "created_at": now_utc,
    })
    mentor_id = str(m_res.inserted_id)
    mentor_user = {"_id": ObjectId(mentor_id), "name": "Prof. Planned Absence Mentor", "role": "MENTOR", "email": "mentor.pa@college.edu"}

    s1_res = await db.students.insert_one({
        "name": "Assigned Student A",
        "roll_number": "24PA001",
        "year": "3rd Year",
        "section": "B",
        "mentor_id": mentor_id,
        "mentor_name": "Prof. Planned Absence Mentor",
        "created_at": now_utc,
        "updated_at": now_utc,
    })
    s1_id = str(s1_res.inserted_id)

    s2_res = await db.students.insert_one({
        "name": "Unassigned Student B",
        "roll_number": "24PA002",
        "year": "3rd Year",
        "section": "B",
        "created_at": now_utc,
        "updated_at": now_utc,
    })
    s2_id = str(s2_res.inserted_id)

    # 2. Mentor attempts to create planned absence for unassigned student -> should fail 403
    with pytest.raises(Exception) as exc_info:
        await create_planned_absence(
            PlannedAbsenceCreateRequest(
                student_id=s2_id,
                roll_number="24PA002",
                start_date="2026-10-10",
                end_date="2026-10-15",
                reason="Medical leave",
            ),
            mentor_user,
        )
    assert "403" in str(exc_info.value) or "Access denied" in str(exc_info.value)

    # 3. Mentor creates planned absence for assigned student -> Success
    pa_rec = await create_planned_absence(
        PlannedAbsenceCreateRequest(
            student_id=s1_id,
            roll_number="24PA001",
            start_date=today_ist,
            end_date="2026-12-31",
            reason="Medical leave for surgery",
        ),
        mentor_user,
    )
    assert pa_rec.id is not None
    assert pa_rec.status == "ACTIVE"
    assert pa_rec.is_active_today is True

    # 4. Overlap detection: Mentor tries to create overlapping absence -> should fail 400
    with pytest.raises(Exception) as exc_overlap:
        await create_planned_absence(
            PlannedAbsenceCreateRequest(
                student_id=s1_id,
                roll_number="24PA001",
                start_date=today_ist,
                end_date="2026-12-25",
                reason="Sports event overlap",
            ),
            mentor_user,
        )
    assert "400" in str(exc_overlap.value) or "already exists" in str(exc_overlap.value)

    # 5. Check active planned absence lookup
    active_pa = await get_active_planned_absence_for_student(s1_id, today_ist)
    assert active_pa is not None
    assert active_pa["reason"] == "Medical leave for surgery"

    # 6. Verify absentee enrichment
    await db.student_attendance.insert_one({
        "date": today_ist,
        "year": "3rd Year",
        "section": "B",
        "roll_number": "24PA001",
        "student_id": s1_id,
        "student_name": "Assigned Student A",
        "status": "Absent",
        "created_at": now_utc,
    })

    absentees = await get_mentor_scoped_absentees("3rd Year", "B", today_ist, mentor_user)
    assert len(absentees) == 1
    assert absentees[0].roll_number == "24PA001"
    assert absentees[0].is_planned_absence is True
    assert absentees[0].planned_absence_reason == "Medical leave for surgery"

    # 7. Update planned absence
    updated_pa = await update_planned_absence(
        pa_rec.id,
        PlannedAbsenceUpdateRequest(
            reason="Updated surgery recovery leave",
        ),
        mentor_user,
    )
    assert updated_pa.reason == "Updated surgery recovery leave"

    # 8. Cancel planned absence
    cancelled_pa = await cancel_planned_absence(
        pa_rec.id,
        PlannedAbsenceCancelRequest(cancellation_reason="Student returned early to college"),
        mentor_user,
    )
    assert cancelled_pa.status == "CANCELLED"
    assert cancelled_pa.is_active_today is False

    # Clean up
    await db.users.delete_one({"_id": ObjectId(mentor_id)})
    await db.students.delete_many({"roll_number": {"$in": ["24PA001", "24PA002"]}})
    await db.student_attendance.delete_many({"roll_number": "24PA001"})
    await db.planned_absences.delete_many({"roll_number": "24PA001"})
