import io
import re
from datetime import datetime, time, timezone
from typing import Any, Dict, List, Tuple
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import EmailStr, TypeAdapter
from app.core.security import hash_password
from app.db.mongodb import get_database
from app.schemas.user import UserRole

email_adapter = TypeAdapter(EmailStr)

DAYS_OF_WEEK = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]


# ----------------------------------------------------
# CR/LR EXCEL PROCESSING
# ----------------------------------------------------

def generate_crlr_excel_template() -> bytes:
    """Generate a clean .xlsx template for CR/LR user import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CR_LR_Import_Template"

    headers = ["Name", "Roll Number", "Email", "Phone", "Role", "Year", "Section"]
    ws.append(headers)

    # Styling header
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sample rows
    samples = [
        ["Rahul Kumar", "22CS2A01", "rahul.kumar@example.com", "9876543210", "CR", "2nd Year", "II-A"],
        ["Sneha Sharma", "22CS2A02", "sneha.sharma@example.com", "9876543211", "LR", "2nd Year", "II-A"],
        ["Vikram Singh", "21CS3B01", "vikram.singh@example.com", "9876543212", "CR", "3rd Year", "III-B"],
    ]
    for row in samples:
        ws.append(row)

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


async def parse_and_import_crlr_excel(
    file_bytes: bytes, filename: str, actor_id: str
) -> Dict[str, Any]:
    """Parse CR/LR Excel workbook and safely import users."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Invalid file format. Only .xlsx Excel workbooks are supported.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ValueError(f"Failed to read Excel file: {str(e)}")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty.")

    header_row = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    expected_cols = ["name", "roll number", "email", "phone", "role", "year", "section"]
    col_map = {}
    for idx, h in enumerate(header_row):
        h_norm = h.lower()
        if h_norm in expected_cols:
            col_map[h_norm] = idx

    missing_cols = [c for c in ["name", "email", "role", "year", "section"] if c not in col_map]
    if missing_cols:
        raise ValueError(f"Missing required columns in Excel: {', '.join(missing_cols)}")

    db = get_database()

    # Pre-fetch existing emails and roll numbers for fast duplicate checking
    existing_users = await db.users.find({}, {"email": 1, "roll_number": 1}).to_list(length=10000)
    existing_emails = {u["email"].lower(): str(u["_id"]) for u in existing_users if "email" in u}
    existing_rolls = {u["roll_number"].upper(): str(u["_id"]) for u in existing_users if u.get("roll_number")}

    seen_emails_in_file = set()
    seen_rolls_in_file = set()

    total_rows = 0
    created_count = 0
    updated_count = 0
    failed_count = 0
    errors = []

    default_password_hash = hash_password("demo1234")

    for row_idx, row_values in enumerate(rows[1:], start=2):
        if not any(row_values):
            continue  # Skip completely empty rows

        total_rows += 1
        row_str_num = f"Row {row_idx}"

        def get_val(col_name: str) -> str:
            idx = col_map.get(col_name)
            if idx is not None and idx < len(row_values) and row_values[idx] is not None:
                return str(row_values[idx]).strip()
            return ""

        name = get_val("name")
        email = get_val("email").lower()
        role = get_val("role").upper()
        roll = get_val("roll number").upper()
        phone = get_val("phone")
        year = get_val("year")
        section = get_val("section").upper()

        row_errors = []

        if not name:
            row_errors.append("Name is required.")
        if not email:
            row_errors.append("Email is required.")
        else:
            try:
                email_adapter.validate_python(email)
            except Exception:
                row_errors.append(f"Invalid email format: '{email}'.")

        if role not in [UserRole.CR.value, UserRole.LR.value]:
            row_errors.append(f"Role must be CR or LR (got '{role}').")

        if not year:
            row_errors.append("Year is required.")
        if not section:
            row_errors.append("Section is required.")

        # Duplicate checks in file
        if email in seen_emails_in_file:
            row_errors.append(f"Duplicate email '{email}' within Excel file.")
        else:
            if email:
                seen_emails_in_file.add(email)

        if roll:
            if roll in seen_rolls_in_file:
                row_errors.append(f"Duplicate roll number '{roll}' within Excel file.")
            else:
                seen_rolls_in_file.add(roll)

        if row_errors:
            failed_count += 1
            errors.append({"row": row_idx, "errors": row_errors})
            continue

        now = datetime.now(timezone.utc)

        # Check if updating or creating
        if email in existing_emails:
            # Update existing user
            u_id = existing_emails[email]
            update_fields = {
                "name": name,
                "role": role,
                "year": year,
                "section": section,
                "updated_at": now,
            }
            if roll:
                update_fields["roll_number"] = roll
            if phone:
                update_fields["phone"] = phone

            await db.users.update_one({"_id": db.users.find_one({"email": email})["_id"]}, {"$set": update_fields})
            updated_count += 1
        elif roll and roll in existing_rolls:
            # Update existing user by roll number
            update_fields = {
                "name": name,
                "email": email,
                "role": role,
                "year": year,
                "section": section,
                "updated_at": now,
            }
            if phone:
                update_fields["phone"] = phone
            await db.users.update_one({"roll_number": roll}, {"$set": update_fields})
            updated_count += 1
        else:
            # Create new user
            new_user_doc = {
                "name": name,
                "email": email,
                "password_hash": default_password_hash,
                "role": role,
                "roll_number": roll if roll else None,
                "phone": phone if phone else None,
                "year": year,
                "section": section,
                "is_active": True,
                "created_at": now,
                "updated_at": now,
            }
            res = await db.users.insert_one(new_user_doc)
            existing_emails[email] = str(res.inserted_id)
            if roll:
                existing_rolls[roll] = str(res.inserted_id)
            created_count += 1

    # Log audit entry
    await db.audit_logs.insert_one(
        {
            "actor_id": actor_id,
            "action": "IMPORT_CRLR_EXCEL",
            "metadata": {
                "total_rows": total_rows,
                "created": created_count,
                "updated": updated_count,
                "failed": failed_count,
            },
            "created_at": datetime.now(timezone.utc),
        }
    )

    return {
        "total_rows": total_rows,
        "created": created_count,
        "updated": updated_count,
        "failed": failed_count,
        "errors": errors,
    }


# ----------------------------------------------------
# TIMETABLE EXCEL PROCESSING
# ----------------------------------------------------

def generate_timetable_excel_template() -> bytes:
    """Generate a clean .xlsx template for Timetable import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class_Timetables"

    headers = [
        "Year",
        "Section",
        "Day",
        "Period",
        "Start Time",
        "End Time",
        "Subject",
        "Faculty",
        "Room",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        ["2nd Year", "II-A", "Monday", 1, "09:10", "10:00", "Database Management Systems", "Dr. A. Sharma", "B-204"],
        ["2nd Year", "II-A", "Monday", 2, "10:00", "10:50", "Database Management Systems", "Dr. A. Sharma", "B-204"],
        ["2nd Year", "II-A", "Monday", 3, "10:50", "11:40", "Database Management Systems", "Dr. A. Sharma", "B-204"],
        ["2nd Year", "II-A", "Monday", 4, "11:40", "12:30", "Operating Systems", "Prof. R. Verma", "B-205"],
        ["2nd Year", "II-B", "Monday", 1, "09:10", "10:00", "Computer Networks", "Dr. K. Patel", "C-101"],
        ["3rd Year", "III-A", "Monday", 1, "09:10", "10:00", "Software Engineering", "Prof. M. Iyer", "A-302"],
    ]
    for row in samples:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _format_time_str(val: Any) -> str:
    """Helper to convert Excel time object or string to HH:MM format."""
    if isinstance(val, (datetime, time)):
        return val.strftime("%H:%M")
    val_str = str(val).strip()
    match = re.match(r"^(\d{1,2}):(\d{2})", val_str)
    if match:
        h, m = int(match.group(1)), int(match.group(2))
        return f"{h:02d}:{m:02d}"
    return val_str


def _get_cell_value(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int) -> Any:
    """Safely get cell value, handling openpyxl MergedCell instances by fetching top-left value."""
    cell = ws.cell(row=row, column=col)
    val = getattr(cell, "value", None)
    if val is not None:
        return val
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return None


def _parse_time_range(time_str: str, period_num: int) -> Tuple[str, str]:
    """Parse time string like '09:10 - 10:00' or '01:30 - 02:20' into 12-hour HH:MM format."""
    parts = re.split(r"\s*[-–toTO]\s*", str(time_str).strip())
    if len(parts) != 2:
        return "", ""

    def convert_time(t_raw: str) -> str:
        t_clean = str(t_raw).strip()
        m = re.match(r"^(\d{1,2}):(\d{2})", t_clean)
        if not m:
            return ""
        h, m_val = int(m.group(1)), int(m.group(2))
        if h > 12:
            h -= 12
        return f"{h:02d}:{m_val:02d}"

    return convert_time(parts[0]), convert_time(parts[1])


def _parse_matrix_timetable_excel(ws: openpyxl.worksheet.worksheet.Worksheet, sheet_name: str) -> List[Dict[str, Any]]:
    """Parse College Matrix Grid Timetable sheet (e.g. DAY | 1 | 2 | Break | 3 | 4 | Lunch | 5 | 6 | 7)."""
    # 1. Search top 12 rows for Section, Year, Semester
    detected_section = None
    detected_year = "2nd Year"

    for r in range(1, min(12, ws.max_row + 1)):
        row_str = " ".join(str(_get_cell_value(ws, r, c) or "") for c in range(1, ws.max_column + 1))
        sec_match = re.search(r"\[\s*([A-Za-z0-9]+\s*-\s*[A-Za-z0-9]+)\s*\]", row_str, re.IGNORECASE)
        if not sec_match:
            sec_match = re.search(r"\b([A-Z]{2,5}\s*-\s*[A-Z0-9]+)\b", row_str)
        if sec_match and not detected_section:
            detected_section = re.sub(r"\s+", "", sec_match.group(1)).upper()

        year_match = re.search(r"\b(I|II|III|IV)\s*B\.?Tech\b", row_str, re.IGNORECASE)
        if year_match:
            roman = year_match.group(1).upper()
            mapping = {"I": "1st Year", "II": "2nd Year", "III": "3rd Year", "IV": "4th Year"}
            detected_year = mapping.get(roman, f"{roman} Year")

    roman_prefix_map = {"1st Year": "I", "2nd Year": "II", "3rd Year": "III", "4th Year": "IV"}
    prefix = roman_prefix_map.get(detected_year, "II")

    if detected_section:
        clean_sec = re.sub(r"^(I|II|III|IV)-", "", detected_section)
        detected_section = f"{prefix}-{clean_sec}"
    else:
        detected_section = f"{prefix}-CSE-A"

    # 2. Locate header row with DAY and Period numbers
    header_row_idx = None
    time_row_idx = None

    for r in range(1, min(15, ws.max_row + 1)):
        for col_check in range(1, min(5, ws.max_column + 1)):
            cell_val = str(_get_cell_value(ws, r, col_check) or "").replace("\n", "").replace(" ", "").upper().strip()
            if "DAY" in cell_val or "PERIOD" in cell_val:
                header_row_idx = r
                time_row_idx = r + 1
                break
        if header_row_idx:
            break

    if not header_row_idx:
        for r in range(1, min(15, ws.max_row + 1)):
            cols_with_digits = 0
            for c in range(1, min(10, ws.max_column + 1)):
                c_val = str(_get_cell_value(ws, r, c) or "").strip()
                if re.search(r"^\d+$", c_val) or "09:" in c_val or "10:" in c_val:
                    cols_with_digits += 1
            if cols_with_digits >= 3:
                header_row_idx = r
                time_row_idx = r + 1
                break

    if not header_row_idx:
        return []

    # Map column indexes to period numbers & time slots
    period_col_map: Dict[int, Tuple[int, str, str]] = {}
    
    for c in range(2, ws.max_column + 1):
        p_val = str(_get_cell_value(ws, header_row_idx, c) or "").strip()
        time_val = str(_get_cell_value(ws, time_row_idx, c) or "").strip()

        if p_val.upper() in ["BREAK", "LUNCH"] or time_val.upper() in ["BREAK", "LUNCH"]:
            continue

        p_num_match = re.search(r"\b(\d+)\b", p_val)
        if not p_num_match:
            p_num_match = re.search(r"\b(\d+)\b", time_val)

        if p_num_match:
            p_num = int(p_num_match.group(1))
        else:
            p_num = len(period_col_map) + 1

        start_t, end_t = _parse_time_range(time_val, p_num)
        if not start_t:
            start_t, end_t = _parse_time_range(p_val, p_num)
        if not start_t:
            default_times = {
                1: ("09:10", "10:00"), 2: ("10:00", "10:50"), 3: ("11:00", "11:50"),
                4: ("11:50", "12:40"), 5: ("13:30", "14:20"), 6: ("14:20", "15:10"), 7: ("15:10", "16:00")
            }
            start_t, end_t = default_times.get(p_num, ("09:00", "10:00"))
        period_col_map[c] = (p_num, start_t, end_t)

    # 3. Parse Day rows (MON to SAT)
    day_mapping = {
        "MON": "Monday", "MONDAY": "Monday",
        "TUE": "Tuesday", "TUESDAY": "Tuesday",
        "WED": "Wednesday", "WEDNESDAY": "Wednesday",
        "THU": "Thursday", "THURSDAY": "Thursday",
        "FRI": "Friday", "FRIDAY": "Friday",
        "SAT": "Saturday", "SATURDAY": "Saturday"
    }

    records = []
    last_day_row_idx = time_row_idx

    for r in range(time_row_idx + 1, ws.max_row + 1):
        raw_day_cell = ""
        for day_c in range(1, min(3, ws.max_column + 1)):
            c_val = str(_get_cell_value(ws, r, day_c) or "").replace("\n", "").replace("\r", "").replace(" ", "").upper().strip()
            if any(c_val.startswith(k) for k in day_mapping):
                raw_day_cell = c_val
                break

        matched_day = None
        for k, d in day_mapping.items():
            if raw_day_cell.startswith(k):
                matched_day = d
                break

        if matched_day:
            last_day_row_idx = r
            for col_idx, (p_num, start_t, end_t) in period_col_map.items():
                cell_raw = str(_get_cell_value(ws, r, col_idx) or "").strip()
                if not cell_raw or cell_raw.upper() in ["BREAK", "LUNCH", "FREE", "NONE"]:
                    continue

                room = None
                cell_no_room = cell_raw
                m_room = re.search(r"\(([^)]+)\)", cell_raw)
                if m_room:
                    room = m_room.group(1).strip()
                    cell_no_room = re.sub(r"\(([^)]+)\)", "", cell_raw).strip()

                faculty_in_cell = None
                lines = [l.strip() for l in re.split(r"[\r\n]+", cell_no_room) if l.strip()]
                subject = lines[0] if lines else cell_no_room
                if len(lines) >= 2:
                    possible_fac = " ".join(lines[1:]).strip()
                    if possible_fac and not re.match(r"^[\d\s\-_/]+$", possible_fac):
                        faculty_in_cell = possible_fac

                records.append({
                    "year": detected_year,
                    "section": detected_section,
                    "day": matched_day,
                    "period": p_num,
                    "start_time": start_t,
                    "end_time": end_t,
                    "subject": subject,
                    "faculty": faculty_in_cell,
                    "room": room,
                    "updated_at": datetime.now(timezone.utc),
                })

    # 4. Parse Faculty Legend Table below last day row (and adjacent cells)
    faculty_legend: Dict[str, str] = {}
    for r in range(last_day_row_idx + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell_val = str(_get_cell_value(ws, r, c) or "").strip()
            if not cell_val:
                continue

            # Case A: Separator in single cell (e.g. "DMGT : Dr. Ramesh", "PYTHON LAB - Prof. Sharma")
            m_sep = re.split(r"\s*[:\-\u2013\u2014]\s*", cell_val, maxsplit=1)
            if len(m_sep) == 2 and m_sep[0].strip() and m_sep[1].strip():
                subj_code = m_sep[0].replace("\n", " ").strip().upper()
                fac_name = m_sep[1].replace("\n", " ").strip()
                if len(fac_name) > 2 and subj_code not in faculty_legend:
                    faculty_legend[subj_code] = fac_name
                continue

            # Case B: Two adjacent cells in a row (Column A = "DMGT", Column B = "Dr. Ramesh")
            if c < ws.max_column:
                adj_val = str(_get_cell_value(ws, r, c + 1) or "").strip()
                if cell_val and adj_val:
                    code_norm = cell_val.replace("\n", " ").strip().upper()
                    name_norm = adj_val.replace("\n", " ").strip()
                    if len(code_norm) <= 25 and len(name_norm) > 2 and code_norm not in faculty_legend:
                        if code_norm not in ["SUBJECT", "COURSE", "SL.NO", "CODE", "PERIOD"]:
                            faculty_legend[code_norm] = name_norm

    # Assign faculty names to matching subject records
    for rec in records:
        if rec["faculty"]:
            continue
        subj_upper = rec["subject"].upper()
        if subj_upper in faculty_legend:
            rec["faculty"] = faculty_legend[subj_upper]
        else:
            for k, v in faculty_legend.items():
                if k in subj_upper or subj_upper in k:
                    rec["faculty"] = v
                    break

    return records


async def parse_and_import_timetable_excel(
    file_bytes: bytes, filename: str, actor_id: str
) -> Dict[str, Any]:
    """Parse timetable workbook (supporting Matrix Grid or Tabular format) and store in MongoDB."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Invalid file format. Only .xlsx files are supported.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ValueError(f"Failed to read Excel workbook: {str(e)}")

    records_to_insert = []
    total_rows = 0
    failed_rows = 0
    errors = []

    # Process all sheets in workbook
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Check if sheet is a College Matrix Grid sheet (e.g. contains DAY in top rows)
        is_matrix_grid = False
        for r in range(1, min(15, ws.max_row + 1)):
            for c in range(1, min(6, ws.max_column + 1)):
                val_upper = str(_get_cell_value(ws, r, c) or "").strip().upper()
                if "DAY" in val_upper or "TIME TABLE" in val_upper or any(k in val_upper for k in ["MON", "TUE", "WED", "THU", "FRI", "SAT"]):
                    is_matrix_grid = True
                    break
            if is_matrix_grid:
                break

        if is_matrix_grid:
            matrix_records = _parse_matrix_timetable_excel(ws, sheet_name)
            records_to_insert.extend(matrix_records)
            total_rows += len(matrix_records)
            continue

        # Standard Tabular parsing fallback
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header_row = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        header_lower = [h.lower() for h in header_row]

        col_map = {}
        for idx, h in enumerate(header_lower):
            col_map[h] = idx

        default_section = sheet_name.strip().upper() if "-" in sheet_name or len(sheet_name) <= 10 else None

        for row_idx, row_values in enumerate(rows[1:], start=2):
            if not any(row_values):
                continue

            total_rows += 1
            row_errors = []

            def get_cell(key: str) -> str:
                idx = col_map.get(key)
                if idx is not None and idx < len(row_values) and row_values[idx] is not None:
                    return str(row_values[idx]).strip()
                return ""

            year = get_cell("year") or "Standard"
            section = get_cell("section").upper() or default_section
            if section and not re.match(r"^(I|II|III|IV)-", section):
                section = f"II-{section}"
            day = get_cell("day").capitalize()
            period_raw = get_cell("period")
            start_time_raw = get_cell("start time") or get_cell("starttime") or get_cell("start_time")
            end_time_raw = get_cell("end time") or get_cell("endtime") or get_cell("end_time")
            subject = get_cell("subject")
            faculty = get_cell("faculty")
            room = get_cell("room")

            if not section:
                row_errors.append("Section name is required.")
            if day not in DAYS_OF_WEEK:
                row_errors.append(f"Invalid Day: '{day}'. Must be one of {DAYS_OF_WEEK}.")

            try:
                period_num = int(re.sub(r"\D", "", period_raw))
            except Exception:
                period_num = 1
                if not period_raw:
                    row_errors.append("Period number is required.")

            start_time = _format_time_str(start_time_raw)
            end_time = _format_time_str(end_time_raw)

            if not re.match(r"^\d{2}:\d{2}$", start_time):
                row_errors.append(f"Invalid Start Time format: '{start_time_raw}'. Expected HH:MM.")
            if not re.match(r"^\d{2}:\d{2}$", end_time):
                row_errors.append(f"Invalid End Time format: '{end_time_raw}'. Expected HH:MM.")
            if not subject:
                row_errors.append("Subject name is required.")

            if row_errors:
                failed_rows += 1
                errors.append({"sheet": sheet_name, "row": row_idx, "errors": row_errors})
                continue

            records_to_insert.append(
                {
                    "year": year,
                    "section": section,
                    "day": day,
                    "period": period_num,
                    "start_time": start_time,
                    "end_time": end_time,
                    "subject": subject,
                    "faculty": faculty if faculty else None,
                    "room": room if room else None,
                    "updated_at": datetime.now(timezone.utc),
                }
            )

    if not records_to_insert and errors:
        return {
            "total_rows": total_rows,
            "inserted": 0,
            "failed": failed_rows,
            "errors": errors,
        }

    db = get_database()

    inserted_count = 0
    section_map = {}
    for rec in records_to_insert:
        await db.timetables.update_one(
            {
                "section": rec["section"],
                "day": rec["day"],
                "period": rec["period"],
            },
            {"$set": rec},
            upsert=True,
        )
        inserted_count += 1
        sec = rec["section"]
        yr = rec.get("year", "2nd Year")
        if sec not in section_map:
            section_map[sec] = yr

    # Auto-upsert sections into db.sections so sessions can be generated
    for sec, yr in section_map.items():
        await db.sections.update_one(
            {"section_name": sec},
            {
                "$set": {
                    "section_name": sec,
                    "year": yr,
                    "is_active": True,
                    "updated_at": datetime.now(timezone.utc),
                },
                "$setOnInsert": {
                    "created_at": datetime.now(timezone.utc),
                    "assigned_cr_id": None,
                    "assigned_lr_id": None,
                },
            },
            upsert=True,
        )

    await db.audit_logs.insert_one(
        {
            "actor_id": actor_id,
            "action": "IMPORT_TIMETABLE_EXCEL",
            "metadata": {
                "total_rows": total_rows,
                "inserted": inserted_count,
                "failed": failed_rows,
            },
            "created_at": datetime.now(timezone.utc),
        }
    )

    return {
        "total_rows": total_rows,
        "inserted": inserted_count,
        "failed": failed_rows,
        "errors": errors,
    }


# ----------------------------------------------------
# MENTOR EXCEL PROCESSING
# ----------------------------------------------------

def generate_mentor_excel_template() -> bytes:
    """Generate a clean .xlsx template for Mentor bulk import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mentor_Import_Template"

    headers = ["Mentor ID", "Mentor Name", "Password", "Phone Number"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        ["M101", "Dr. A. Ramesh", "mentor1234", "9876543210"],
        ["M102", "Prof. S. Sunita", "mentor1234", "9876543211"],
    ]
    for row in samples:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


async def parse_and_import_mentor_excel(
    file_bytes: bytes, filename: str, actor_id: str
) -> Dict[str, Any]:
    """Parse Mentor Excel workbook and safely import mentors into users collection."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Invalid file format. Only .xlsx Excel workbooks are supported.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ValueError(f"Failed to read Excel file: {str(e)}")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty.")

    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    col_map = {}
    for idx, h in enumerate(header_row):
        if "mentor id" in h or h == "id":
            col_map["mentor_id"] = idx
        elif "name" in h:
            col_map["name"] = idx
        elif "password" in h:
            col_map["password"] = idx
        elif "phone" in h:
            col_map["phone"] = idx

    missing_cols = [c for c in ["mentor_id", "name"] if c not in col_map]
    if missing_cols:
        raise ValueError(f"Missing required columns in Excel: {', '.join(missing_cols)}")

    db = get_database()
    existing_users = await db.users.find({"role": UserRole.MENTOR.value}).to_list(length=5000)
    existing_mentors = {u["mentor_id"].upper(): str(u["_id"]) for u in existing_users if u.get("mentor_id")}

    seen_ids_in_file = set()
    total_rows = 0
    created_count = 0
    updated_count = 0
    failed_count = 0
    errors = []

    for row_idx, row_values in enumerate(rows[1:], start=2):
        if not any(row_values):
            continue

        total_rows += 1

        def get_val(col_name: str) -> str:
            idx = col_map.get(col_name)
            if idx is not None and idx < len(row_values) and row_values[idx] is not None:
                return str(row_values[idx]).strip()
            return ""

        mentor_id = get_val("mentor_id").upper()
        name = get_val("name")
        password = get_val("password") or "mentor1234"
        phone = get_val("phone")

        row_errors = []
        if not mentor_id:
            row_errors.append("Mentor ID is required.")
        if not name:
            row_errors.append("Mentor Name is required.")

        if mentor_id in seen_ids_in_file:
            row_errors.append(f"Duplicate Mentor ID '{mentor_id}' within Excel file.")
        else:
            if mentor_id:
                seen_ids_in_file.add(mentor_id)

        if row_errors:
            failed_count += 1
            errors.append({"row": row_idx, "mentor_id": mentor_id or "N/A", "errors": row_errors})
            continue

        now = datetime.now(timezone.utc)
        email = f"{mentor_id.lower()}@fams.edu"

        if mentor_id in existing_mentors:
            u_id = existing_mentors[mentor_id]
            update_data = {"name": name, "updated_at": now}
            if phone:
                update_data["phone"] = phone
            if password:
                update_data["password_hash"] = hash_password(password)
            await db.users.update_one({"_id": ObjectId(u_id)}, {"$set": update_data})
            updated_count += 1
        else:
            new_doc = {
                "name": name,
                "email": email,
                "mentor_id": mentor_id,
                "password_hash": hash_password(password),
                "role": UserRole.MENTOR.value,
                "phone": phone if phone else None,
                "is_active": True,
                "created_at": now,
                "updated_at": now,
            }
            res = await db.users.insert_one(new_doc)
            existing_mentors[mentor_id] = str(res.inserted_id)
            created_count += 1

    await db.audit_logs.insert_one({
        "actor_id": actor_id,
        "action": "IMPORT_MENTORS_EXCEL",
        "metadata": {"total_rows": total_rows, "created": created_count, "updated": updated_count, "failed": failed_count},
        "created_at": datetime.now(timezone.utc),
    })

    return {
        "total_rows": total_rows,
        "created": created_count,
        "updated": updated_count,
        "failed": failed_count,
        "errors": errors,
    }


# ----------------------------------------------------
# STUDENT EXCEL PROCESSING
# ----------------------------------------------------

def generate_student_excel_template() -> bytes:
    """Generate a clean .xlsx template for Student bulk import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student_Import_Template"

    headers = ["Year", "Name", "Roll Number", "Section", "Student Phone Number", "Parent Phone Number"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        ["2nd Year", "Rahul Kumar", "23CS001", "CSE-A", "9876543210", "9876543299"],
        ["2nd Year", "Sneha Sharma", "23CS002", "CSE-A", "9876543211", "9876543298"],
        ["3rd Year", "Vikram Singh", "22CS010", "CSE-B", "9876543212", "9876543297"],
    ]
    for row in samples:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


async def parse_and_import_student_excel(
    file_bytes: bytes, filename: str, actor_id: str
) -> Dict[str, Any]:
    """Parse Student Excel workbook and safely import students into students collection."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Invalid file format. Only .xlsx Excel workbooks are supported.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ValueError(f"Failed to read Excel file: {str(e)}")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty.")

    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    col_map = {}
    for idx, h in enumerate(header_row):
        if "year" in h:
            col_map["year"] = idx
        elif "roll" in h:
            col_map["roll_number"] = idx
        elif "section" in h:
            col_map["section"] = idx
        elif "parent" in h:
            col_map["parent_phone"] = idx
        elif "student phone" in h or (("phone" in h) and "parent" not in h):
            col_map["student_phone"] = idx
        elif "name" in h:
            col_map["name"] = idx

    missing_cols = [c for c in ["year", "name", "roll_number", "section"] if c not in col_map]
    if missing_cols:
        raise ValueError(f"Missing required columns in Excel: {', '.join(missing_cols)}")

    db = get_database()
    existing_students = await db.students.find({}, {"roll_number": 1}).to_list(length=10000)
    existing_rolls = {s["roll_number"].upper(): str(s["_id"]) for s in existing_students if s.get("roll_number")}

    seen_rolls_in_file = set()
    total_rows = 0
    created_count = 0
    updated_count = 0
    failed_count = 0
    errors = []

    for row_idx, row_values in enumerate(rows[1:], start=2):
        if not any(row_values):
            continue

        total_rows += 1

        def get_val(col_name: str) -> str:
            idx = col_map.get(col_name)
            if idx is not None and idx < len(row_values) and row_values[idx] is not None:
                return str(row_values[idx]).strip()
            return ""

        year = get_val("year")
        name = get_val("name")
        roll = get_val("roll_number").upper()
        section = get_val("section").upper()
        student_phone = get_val("student_phone")
        parent_phone = get_val("parent_phone")

        row_errors = []
        if not year:
            row_errors.append("Year is required.")
        if not name:
            row_errors.append("Student Name is required.")
        if not roll:
            row_errors.append("Roll Number is required.")
        if not section:
            row_errors.append("Section is required.")

        if roll in seen_rolls_in_file:
            row_errors.append(f"Duplicate roll number '{roll}' within Excel file.")
        else:
            if roll:
                seen_rolls_in_file.add(roll)

        if row_errors:
            failed_count += 1
            errors.append({"row": row_idx, "roll_number": roll or "N/A", "errors": row_errors})
            continue

        now = datetime.now(timezone.utc)
        student_doc = {
            "year": year,
            "name": name,
            "roll_number": roll,
            "section": section,
            "student_phone": student_phone if student_phone else None,
            "parent_phone": parent_phone if parent_phone else None,
            "updated_at": now,
        }

        if roll in existing_rolls:
            s_id = existing_rolls[roll]
            await db.students.update_one({"_id": ObjectId(s_id)}, {"$set": student_doc})
            updated_count += 1
        else:
            student_doc["created_at"] = now
            res = await db.students.insert_one(student_doc)
            existing_rolls[roll] = str(res.inserted_id)
            created_count += 1

    await db.audit_logs.insert_one({
        "actor_id": actor_id,
        "action": "IMPORT_STUDENTS_EXCEL",
        "metadata": {"total_rows": total_rows, "created": created_count, "updated": updated_count, "failed": failed_count},
        "created_at": datetime.now(timezone.utc),
    })

    return {
        "total_rows": total_rows,
        "created": created_count,
        "updated": updated_count,
        "failed": failed_count,
        "errors": errors,
    }
